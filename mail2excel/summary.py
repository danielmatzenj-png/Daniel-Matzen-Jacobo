"""Resumen diario: compone el texto y lo envía por correo (Apple Mail)."""

from __future__ import annotations

import platform
import subprocess
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .excel_writer import HEADER_TO_KEY


def read_rows(path: str | Path, sheet_name: str) -> list[dict[str, Any]]:
    """Lee la tabla y devuelve una lista de dicts con claves internas."""
    p = Path(path)
    if not p.exists():
        return []
    wb = load_workbook(p, read_only=True)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = next(rows_iter)
    except StopIteration:
        return []
    keys = [HEADER_TO_KEY.get(str(h).strip().lower(), str(h).strip().lower()) for h in headers]

    rows: list[dict[str, Any]] = []
    for row in rows_iter:
        if row is None or all(c is None for c in row):
            continue
        rows.append(
            {keys[i]: ("" if (i >= len(row) or row[i] is None) else row[i]) for i in range(len(keys))}
        )
    return rows


def build_summary(rows: list[dict[str, Any]], date_label: str) -> tuple[str, str]:
    """Devuelve (asunto_sufijo, cuerpo) del resumen para el día indicado."""
    todays = [r for r in rows if str(r.get("day", "")).strip() == date_label]

    lines: list[str] = []
    lines.append(f"Resumen de BLs registrados el {date_label}")
    lines.append("=" * 40)
    lines.append(f"Total de correos registrados hoy: {len(todays)}")
    lines.append("")

    # Conteo por cliente.
    by_customer: dict[str, int] = {}
    for r in todays:
        cust = str(r.get("customer", "") or "—")
        by_customer[cust] = by_customer.get(cust, 0) + 1
    if by_customer:
        lines.append("Por cliente:")
        for cust, n in sorted(by_customer.items(), key=lambda x: (-x[1], x[0])):
            lines.append(f"  · {cust}: {n}")
        lines.append("")

    # Detalle.
    if todays:
        lines.append("Detalle:")
        for r in todays:
            bl = r.get("bl_nr") or "(sin BL)"
            ref = r.get("internal_reference") or "(sin ref)"
            lines.append(
                f"  · {r.get('customer','')} | BL {bl} | Ref {ref} | "
                f"ETD {r.get('etd','') or '-'} | ETA {r.get('eta','') or '-'}"
            )
        lines.append("")

    # Anomalías.
    anomalies = [r for r in todays if str(r.get("notes", "")).strip()]
    if anomalies:
        lines.append(f"⚠ Anomalías ({len(anomalies)}):")
        for r in anomalies:
            lines.append(f"  · {r.get('customer','')} (BL {r.get('bl_nr') or '—'}): {r.get('notes')}")
    else:
        lines.append("Sin anomalías registradas hoy.")

    return date_label, "\n".join(lines)


def _applescript_str(text: str) -> str:
    """Escapa comillas y barras para incrustar texto en un literal AppleScript."""
    return text.replace("\\", "\\\\").replace('"', '\\"')


def send_via_apple_mail(to: str, subject: str, body: str) -> None:
    """Envía un correo desde Apple Mail (solo macOS)."""
    script = f'''
tell application "Mail"
    set newMsg to make new outgoing message with properties {{subject:"{_applescript_str(subject)}", content:"{_applescript_str(body)}", visible:false}}
    tell newMsg
        make new to recipient at end of to recipients with properties {{address:"{_applescript_str(to)}"}}
    end tell
    send newMsg
end tell
'''
    proc = subprocess.run(["osascript", "-"], input=script, capture_output=True, text=True, timeout=120)
    if proc.returncode != 0:
        raise RuntimeError(f"No se pudo enviar el resumen por Apple Mail: {proc.stderr.strip()}")


def send_summary(cfg: dict[str, Any], date_label: str | None = None) -> str:
    """Compone y envía el resumen diario. Devuelve el cuerpo generado."""
    out = cfg.get("output", {})
    fmt = out.get("date_format", "%d.%m.%y")
    date_label = date_label or datetime.now().strftime(fmt)

    rows = read_rows(out.get("path", "Tabla BLs.xlsx"), out.get("sheet", "Sheet1"))
    _, body = build_summary(rows, date_label)

    summ = cfg.get("summary", {})
    to = summ.get("to", "")
    subject = f"{summ.get('subject_prefix', 'Resumen diario BLs')} — {date_label}"

    send_via = summ.get("send_via", "applescript")
    if send_via == "applescript" and platform.system() == "Darwin" and to:
        send_via_apple_mail(to, subject, body)
        print(f"✓ Resumen enviado a {to}")
    else:
        print(f"(No se envió por correo: send_via={send_via}, macOS={platform.system()=='Darwin'})")
        print(f"\n--- {subject} ---\n{body}\n")
    return body
