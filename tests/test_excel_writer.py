"""Pruebas de la escritura a Excel y la deduplicación entre ejecuciones."""

from __future__ import annotations

from openpyxl import load_workbook

from mail2excel.config import field_order, load_config
from mail2excel.excel_writer import write_rows
from mail2excel.extractor import extract_all
from mail2excel.mail_reader import read_from_json


def _rows():
    cfg = load_config()
    emails = read_from_json("tests/sample_emails.json")
    return extract_all(emails, fields=cfg["fields"], filters=cfg["filters"]), field_order(cfg)


def test_write_creates_file(tmp_path):
    rows, order = _rows()
    out = tmp_path / "facturas.xlsx"

    summary = write_rows(rows, out, order)
    assert out.exists()
    assert summary["nuevas"] == 3
    assert summary["omitidas"] == 0
    assert summary["total"] == 3

    wb = load_workbook(out)
    ws = wb["Facturas"]
    headers = [c.value for c in ws[1]]
    assert headers[:4] == ["fecha_correo", "remitente", "correo_remitente", "asunto"]
    assert "total" in headers


def test_rerun_deduplicates(tmp_path):
    rows, order = _rows()
    out = tmp_path / "facturas.xlsx"

    write_rows(rows, out, order)
    summary = write_rows(rows, out, order)  # segunda ejecución, mismos correos

    assert summary["nuevas"] == 0
    assert summary["omitidas"] == 3
    assert summary["total"] == 3

    wb = load_workbook(out)
    ws = wb["Facturas"]
    assert ws.max_row - 1 == 3  # sin duplicar
